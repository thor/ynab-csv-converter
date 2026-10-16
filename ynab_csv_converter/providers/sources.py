"""Row readers for the file formats banks export.

Each factory returns a `RowReader`. Blank rows are skipped everywhere.
"""

import csv
from collections.abc import Iterator, Mapping
from pathlib import Path

from .base import NumberedRow, RowReader


def csv_source(
    *, encoding: str = "utf-8", delimiter: str = ",", skip_rows: int = 0, quotechar: str = '"'
) -> RowReader:
    """Read a CSV file. `skip_rows` drops header lines."""

    def read(path: Path) -> Iterator[NumberedRow]:
        with open(path, encoding=encoding, newline="") as handle:
            reader = csv.reader(handle, delimiter=delimiter, quotechar=quotechar)
            for row in reader:
                if reader.line_num <= skip_rows or _is_blank(row):
                    continue
                yield reader.line_num, row

    return read


def xlsx_source(*, sheet: str | None = None, skip_rows: int = 0) -> RowReader:
    """Read an Excel workbook. Uses the active sheet unless `sheet` is given."""

    def read(path: Path) -> Iterator[NumberedRow]:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = workbook[sheet] if sheet else workbook.active
            for line_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                if line_number <= skip_rows or _is_blank(row):
                    continue
                yield line_number, row
        finally:
            workbook.close()

    return read


def by_suffix(readers: Mapping[str, RowReader]) -> RowReader:
    """Pick a reader from the file suffix, e.g. {".csv": csv_source(), ".xlsx": xlsx_source()}."""

    def read(path: Path) -> Iterator[NumberedRow]:
        suffix = path.suffix.lower()
        if suffix not in readers:
            supported = ", ".join(sorted(readers))
            raise ValueError(f"{path}: unsupported file type {suffix!r}, expected one of {supported}")
        yield from readers[suffix](path)

    return read


def _is_blank(row) -> bool:
    return all(cell is None or str(cell).strip() == "" for cell in row)
